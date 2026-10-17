"""Relevé de mesure d'un deck hérité (lecture seule) — appelant : `python3 scripts/releve_pptx.py <deck.pptx>`, consommé par le §0 de la spec et le chiffrage par deck."""
import io
import math
import re
import statistics
import sys
from collections import Counter

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Emu

NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main",
      "p": "http://schemas.openxmlformats.org/presentationml/2006/main"}
POUCE = 914400
PASTILLE_MAX = Emu(int(0.72 * POUCE))   # côté max d'une pastille numérotée
ZONE_GAP = 0.5 * POUCE                  # pastille sur capture ↔ bord du cadre de zone
CONSIGNE_DIST = 1.6 * POUCE             # pastille hors capture ↔ centre du texte le plus proche
DENSE = 90                              # mots par diapositive
INACHEVE = ["a rédiger", "a supprimer", "insérer", "xxx"]


def formes(shapes, tf=lambda x, y, w, h: (x, y, w, h)):
    """Aplatit l'arbre des formes en (forme, boîte absolue en EMU), groupes compris."""
    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            xf = sh._element.find("p:grpSpPr/a:xfrm", NS)
            o, e = xf.find("a:off", NS), xf.find("a:ext", NS)
            co, ce = xf.find("a:chOff", NS), xf.find("a:chExt", NS)
            ox, oy, ex, ey = (int(o.get("x")), int(o.get("y")), int(e.get("cx")), int(e.get("cy")))
            cx, cy = int(co.get("x")), int(co.get("y"))
            cex, cey = int(ce.get("cx")) or 1, int(ce.get("cy")) or 1

            def ntf(x, y, w, h, tf=tf, ox=ox, oy=oy, ex=ex, ey=ey, cx=cx, cy=cy, cex=cex, cey=cey):
                sx, sy = ex / cex, ey / cey
                return tf(ox + (x - cx) * sx, oy + (y - cy) * sy, w * sx, h * sy)
            yield from formes(sh.shapes, ntf)
        elif sh.left is not None:
            yield sh, tf(sh.left, sh.top, sh.width, sh.height)


def texte(sh):
    return sh.text_frame.text.strip() if sh.has_text_frame else ""


def est_image(sh):
    if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
        return True
    try:
        sh.image
        return True
    except Exception:
        return False


def centre(b):
    return b[0] + b[2] / 2, b[1] + b[3] / 2


def dedans(b, c):
    x, y = centre(b)
    return c[0] <= x <= c[0] + c[2] and c[1] <= y <= c[1] + c[3]


def ecart_bord(b, r):
    x, y = centre(b)
    dx = max(r[0] - x, 0, x - (r[0] + r[2]))
    dy = max(r[1] - y, 0, y - (r[1] + r[3]))
    return math.hypot(dx, dy)


def releve(chemin):
    prs = Presentation(chemin)
    sw = prs.slide_width
    m = Counter()
    resolutions, denses, inacheves, lignes = [], [], [], []
    for n, s in enumerate(prs.slides, 1):
        items = list(formes(s.shapes))
        caps = [(sh, b) for sh, b in items if est_image(sh) and b[2] > sw * 0.25]
        cadres = [b for _, b in caps]
        pastilles = [(texte(sh), b) for sh, b in items
                     if sh.shape_type in (MSO_SHAPE_TYPE.AUTO_SHAPE, MSO_SHAPE_TYPE.TEXT_BOX)
                     and re.fullmatch(r"\d{1,2}", texte(sh)) and b[2] < PASTILLE_MAX and b[3] < PASTILLE_MAX]
        zones = [b for sh, b in items if sh.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and not texte(sh)
                 and sh._element.find(".//p:spPr/a:noFill", NS) is not None]
        textes = [b for sh, b in items if sh.has_text_frame and len(texte(sh)) > 8]
        notes = s.notes_slide.notes_text_frame.text.strip() if s.has_notes_slide else ""
        mots = sum(len(texte(sh).split()) for sh, _ in items if sh.has_text_frame)
        sur = [b for _, b in pastilles if any(dedans(b, c) for c in cadres)]
        hors = [b for _, b in pastilles if not any(dedans(b, c) for c in cadres)]
        nature = "pas-a-pas" if caps and pastilles else "capture" if caps else "autre"
        m[nature] += 1
        m["notes"] += bool(notes)
        m["mots"] += mots
        if caps:
            m["pastilles_sur"] += len(sur)
            m["pastilles_sur_zone"] += sum(1 for b in sur if zones and min(ecart_bord(b, z) for z in zones) < ZONE_GAP)
            m["pastilles_hors"] += len(hors)
            m["pastilles_hors_texte"] += sum(1 for b in hors if textes and min(
                math.dist(centre(b), centre(t)) for t in textes) < CONSIGNE_DIST)
        for sh, b in caps:
            try:
                largeur = Image.open(io.BytesIO(sh.image.blob)).size[0]
                resolutions.append(largeur / (b[2] / sw * 1920))
            except Exception:
                m["captures_illisibles"] += 1
        if mots > DENSE:
            denses.append((n, mots))
        corpus = (" ".join(texte(sh) for sh, _ in items) + " " + notes).lower()
        inacheves += [(n, p) for p in INACHEVE if p in corpus]
        lignes.append((n, s.slide_layout.name, nature, mots))
    return prs, m, resolutions, denses, inacheves, lignes


def sequences(prs, lignes):
    """Découpe aux dispositions « Slide transition » ; compte les temps par séquence."""
    out, cur = [], {"titre": "(ouverture)", "n0": 1, "l": []}
    for (n, lay, nature, mots), s in zip(lignes, prs.slides):
        titre = " | ".join(texte(sh) for sh in s.shapes if sh.has_text_frame and texte(sh)).replace("\n", " ")
        if "Slide transition" in lay:
            out.append(cur)
            cur = {"titre": titre, "n0": n, "l": []}
        cur["l"].append((lay, nature, mots, titre.lower()))
    out.append(cur)
    return out


def main(chemin):
    prs, m, res, denses, inach, lignes = releve(chemin)
    print(f"# Relevé — {chemin.split('/')[-1]}")
    print(f"diapositives {len(prs.slides)} · dispositions {len({l[1] for l in lignes})} · masques {len(prs.slide_masters)}")
    print(f"mots {m['mots']} · médiane/diapo {statistics.median(l[3] for l in lignes)} · denses (>{DENSE}) {len(denses)}")
    print(f"notes {m['notes']}/{len(prs.slides)}")
    print(f"pas-à-pas {m['pas-a-pas']} · captures sans pastille {m['capture']} · autres {m['autre']}")
    print(f"pastilles sur capture {m['pastilles_sur']}, dont accolées à une zone {m['pastilles_sur_zone']}")
    print(f"pastilles hors capture {m['pastilles_hors']}, dont proches d'un texte {m['pastilles_hors_texte']}")
    if res:
        print(f"résolution des captures (px source / px affiché à 1920) : médiane {statistics.median(res):.2f}, "
              f"sous 1,0 : {sum(r < 1 for r in res)}/{len(res)}")
    print(f"marqueurs d'inachèvement : {inach}")
    print("\nséquences (n0 · diapos · démo · pas-à-pas · captures · à vous de jouer · règles du jeu · retenez · mots)")
    for q in sequences(prs, lignes):
        c = Counter()
        for lay, nature, mots, t in q["l"]:
            c["demo"] += "Ecran Votre promotion" in lay and "démonstration" in t
            c["avj"] += "Ecran Votre promotion" in lay and "a vous de jouer" in t
            c["rdj"] += "Regles du jeu" in lay
            c["ret"] += "Que retenez" in lay
            c[nature] += 1
            c["mots"] += mots
        print(f"{q['n0']:>4} {len(q['l']):>3} {c['demo']:>3} {c['pas-a-pas']:>3} {c['capture']:>3} {c['avj']:>3} "
              f"{c['rdj']:>3} {c['ret']:>3} {c['mots']:>5}  {q['titre'][:56]}")


if __name__ == "__main__":
    main(sys.argv[1])
